import copy, json, os
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SIM = json.load(open(os.path.join(HERE, 'sim.json')))

# ---- palette: template blue + logo gold + product green ----
NAVY, BLUE, BRIGHT, GOLD, GREEN, RED = '14305E', '1F4E8C', '2F6FD0', 'F2B705', '12A37A', 'D64545'
LIGHT, MID, TEXT, MUTED, WHITE, SOFTG, SOFTY, SOFTR = 'EEF3FA', 'C9D6EA', '1B1F2A', '5B6678', 'FFFFFF', 'E6F6F0', 'FFF6D6', 'FBE9E9'
FONT = 'Calibri'
I = Inches
rgb = RGBColor.from_string


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.25, rad=0.1, shape=None, name=None):
    shp = sl.shapes.add_shape(shape or MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    if shape is None: shp.adjustments[0] = min(0.5, rad / min(w, h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    if line: shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    if name: shp.name = name
    return shp


def _bullet(p, indent=0.2, char='•', color=None):
    pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(indent * 914400))); pPr.set('indent', str(-int(indent * 914400)))
    if color:
        bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', color)
    bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
    bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', char)


def fill_tf(tf, paras, size=16, color=TEXT, bold=False, align='l', space=4, font=FONT, bullet=False, lsp=None):
    first = True
    for item in paras:
        o = {'size': size, 'color': color, 'bold': bold, 'align': align, 'space': space, 'bullet': bullet}
        if isinstance(item, str): runs = [(item, {})]
        elif isinstance(item, list): runs = item
        else:
            o.update({k: v for k, v in item.items() if k not in ('t', 'runs')})
            runs = item['runs'] if 'runs' in item else [(item['t'], {})]
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[o['align']]
        p.space_after = Pt(o['space']); p.space_before = Pt(0)
        if lsp: p.line_spacing = lsp
        for t, ro in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(ro.get('size', o['size'])); f.bold = ro.get('bold', o['bold']); f.italic = ro.get('italic', False); f.name = font
            f.color.rgb = rgb(ro.get('color', o['color']))
            if ro.get('u'): f.underline = True
        if o['bullet']: _bullet(p, color=o.get('bucolor'))


def txt(sl, x, y, w, h, paras, anchor='t', margin=0.06, **kw):
    tb = sl.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = I(margin); tf.margin_top = tf.margin_bottom = I(0.03)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, **kw)
    return tb


def box(sl, x, y, w, h, paras, fill=LIGHT, line=None, anchor='m', rad=0.1, margin=0.1, **kw):
    s = rect(sl, x, y, w, h, fill, line, rad=rad)
    tf = s.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = I(margin); tf.margin_top = tf.margin_bottom = I(0.04)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, **kw)
    return s


def panel(sl, x, y, w, h, title, accent=BLUE, hh=0.58):
    rect(sl, x, y, w, h, WHITE, MID, 1.5, rad=0.14)
    rect(sl, x, y, w, hh, accent, rad=0.14)
    rect(sl, x, y + hh / 2, w, hh / 2, accent, shape=MSO_SHAPE.RECTANGLE)
    txt(sl, x + 0.18, y, w - 0.36, hh, [title], size=21, color=WHITE, bold=True, anchor='m')


def arrow(sl, x1, y1, x2, y2, color=MUTED, w=2.5, head=True, dash=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', 'dash')
    if head:
        t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c


def circle(sl, cx, cy, d, fill, text='', size=16, color=WHITE, line=None):
    s = rect(sl, cx - d / 2, cy - d / 2, d, d, fill, line, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if text: fill_tf(tf, [text], size=size, color=color, bold=True, align='c', space=0)
    return s


def style_chart(ch, size=13):
    ch.font.size = Pt(size); ch.font.name = FONT; ch.font.color.rgb = rgb(TEXT)


def bar_chart(sl, x, y, w, h, cats, series, colors, title=None, horizontal=False, size=13, fmt='General', legend=False, gap=60, ymax=None, point_colors=None):
    cd = CategoryChartData(); cd.categories = cats
    for n, v in series: cd.add_series(n, v)
    typ = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = sl.shapes.add_chart(typ, I(x), I(y), I(w), I(h), cd).chart
    style_chart(ch, size)
    ch.has_title = bool(title)
    if title: ch.chart_title.text_frame.text = title; ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size + 1); ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.has_legend = legend
    if legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(size)
    pl = ch.plots[0]; pl.gap_width = gap; pl.has_data_labels = True
    dl = pl.data_labels; dl.font.size = Pt(size); dl.font.bold = True; dl.number_format = fmt; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for s, c in zip(pl.series, colors):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(c)
    if point_colors:
        for i, c in enumerate(point_colors):
            pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
    va = ch.value_axis; va.has_major_gridlines = False; va.visible = False
    if ymax: va.maximum_scale = ymax
    va.minimum_scale = 0 if min(min(v for v in s[1] if v is not None) for s in series) >= 0 else None
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(size); ca.format.line.color.rgb = rgb(MID); ca.has_major_gridlines = False
    return ch


def table(sl, x, y, w, rows, colw, rowh, size=14, head=NAVY, zebra=True, align=None, first_bold=True):
    nr, nc = len(rows), len(rows[0])
    gf = sl.shapes.add_table(nr, nc, I(x), I(y), I(w), I(rowh * nr)); tb = gf.table
    tb.first_row = False; tb.horz_banding = False
    for j, cw in enumerate(colw): tb.columns[j].width = I(cw)
    for i in range(nr): tb.rows[i].height = I(rowh)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tb.cell(i, j); o = val if isinstance(val, dict) else {'t': val}
            cell.margin_left = cell.margin_right = I(0.08); cell.margin_top = cell.margin_bottom = I(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(o.get('fill', head if i == 0 else (LIGHT if zebra and i % 2 == 0 else WHITE)))
            tf = cell.text_frame; tf.word_wrap = True
            a = (align[j] if align else 'l')
            fill_tf(tf, [o['t']], size=o.get('size', size), color=o.get('c', WHITE if i == 0 else TEXT), bold=o.get('b', i == 0 or (first_bold and j == 0)), align=o.get('a', 'c' if (i == 0 and a == 'c') else a), space=0)
    return gf


def src(sl, text):
    txt(sl, 0.45, 10.02, 19.1, 0.36, [text], size=12, color=MUTED, anchor='m')


def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


def chip(sl, x, y, w, h, text, fill=LIGHT, color=BLUE, size=14, line=None, bold=True):
    return box(sl, x, y, w, h, [text], fill=fill, line=line, size=size, color=color, bold=bold, align='c', rad=0.08, margin=0.04, space=0)


def remove_shape(sl, name):
    for sh in list(sl.shapes):
        if sh.name == name: sh._element.getparent().remove(sh._element)


def crop(path, out, aspect):
    im = Image.open(path).convert('RGB'); w, h = im.size
    if w / h > aspect: nw = int(h * aspect); im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
    else: nh = int(w / aspect); im = im.crop((0, 0, w, nh))
    im.save(out, quality=92); return out


